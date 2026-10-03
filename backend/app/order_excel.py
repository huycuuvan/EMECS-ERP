"""File Excel đặt hàng theo mẫu của khách — "BẢNG XÁC NHẬN GIÁ TRỊ VÀ KHỐI LƯỢNG ĐẶT HÀNG".

Cột chuẩn: STT · Tên hàng hóa · ĐVT · Số lượng · KL/1 bộ · Tổng KL (kg) · Đơn giá · Thành tiền · Ghi chú;
cuối bảng: TỔNG CỘNG TRƯỚC THUẾ · THUẾ VAT x% · TỔNG CỘNG SAU THUẾ.

- `parse_order_excel`: đọc file khách gửi → dòng hàng để điền vào form Đơn hàng (chưa lưu). Tìm hàng tiêu đề theo
  tên cột (không phụ thuộc vị trí), tự tính lại Tổng KL = SL × KL/1 bộ và Thành tiền = Tổng KL × Đơn giá.
- `order_excel`: xuất đơn hàng ra đúng mẫu đó (đơn rỗng = file mẫu để tải về).
"""
import re
import unicodedata
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITLE = "BẢNG XÁC NHẬN GIÁ TRỊ VÀ KHỐI LƯỢNG ĐẶT HÀNG"
HEADERS = ["STT", "Tên hàng hóa", "ĐVT", "Số lượng", "KL/1 bộ", "Tổng KL (kg)", "Đơn giá", "Thành tiền", "Ghi chú"]
CUSTOMER_PLACEHOLDER = "(TỰ LẤY THÔNG TIN KHÁCH HÀNG)"


def _norm(v) -> str:
    """Bỏ dấu, chữ thường, gom khoảng trắng — so khớp tên cột / nhãn."""
    s = unicodedata.normalize("NFD", str(v or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn").lower()
    return re.sub(r"\s+", " ", s).strip()


# khóa → các cách viết tiêu đề cột (đã bỏ dấu); thứ tự quan trọng: "tong kl" phải khớp trước "kl"
COLUMNS = [
    ("stt", ("stt", "tt")),
    ("total_kg", ("tong kl", "tong khoi luong", "kl tong")),
    ("kg_per_unit", ("kl/1", "kl / 1", "kl 1 ", "khoi luong/1", "khoi luong 1", "kl/bo", "kl/ 1")),
    ("name", ("ten hang", "hang muc", "ten san pham", "noi dung")),
    ("unit", ("dvt", "don vi")),
    ("qty", ("so luong", "sl")),
    ("price", ("don gia",)),
    ("amount", ("thanh tien",)),
    ("note", ("ghi chu",)),
]


def _match(text: str) -> str | None:
    t = _norm(text) + " "
    for key, keys in COLUMNS:
        if any(t == k.strip() + " " or t.startswith(k) for k in keys):
            return key
    return None


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("₫", "").replace("đ", "")
    if not s or s.startswith("="):
        return None
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s):  # 1.234.567,5 (kiểu Việt Nam)
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", s):  # 1,234,567.5
        s = s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_order_excel(data: bytes, filename: str = "") -> dict:
    if filename and not filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Chỉ đọc được file .xlsx — mở file bằng Excel và Lưu thành (Save As) dạng .xlsx")
    try:
        wb_val = load_workbook(BytesIO(data), data_only=True)  # giá trị đã tính của công thức (nếu Excel đã lưu)
    except Exception:
        raise HTTPException(400, "Không đọc được file Excel — kiểm tra lại file (.xlsx)")

    for ws in wb_val.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        header_idx, cols = None, {}
        for r, row in enumerate(rows[:40]):
            found = {}
            for c, v in enumerate(row):
                k = _match(v) if isinstance(v, str) else None
                if k and k not in found:
                    found[k] = c
            if "name" in found and ("qty" in found or "total_kg" in found) and "price" in found:
                header_idx, cols = r, found
                break
        if header_idx is None:
            continue
        return _read_sheet(ws.title, rows, header_idx, cols)
    raise HTTPException(400, "Không tìm thấy bảng hàng hóa — file cần có các cột: " + ", ".join(HEADERS[1:8]))


def _read_sheet(sheet: str, rows: list[list], header_idx: int, cols: dict) -> dict:
    warnings: list[str] = []
    get = lambda row, k: row[cols[k]] if k in cols and cols[k] < len(row) else None  # noqa: E731

    # công ty đặt hàng: ô bên phải nhãn "CÔNG TY ĐẶT HÀNG"
    customer = ""
    for row in rows[:header_idx]:
        for c, v in enumerate(row):
            if isinstance(v, str) and _norm(v).startswith("cong ty dat hang"):
                rest = [x for x in row[c + 1:] if x not in (None, "")]
                inline = v.split(":", 1)[1].strip() if ":" in v else ""
                cand = str(rest[0]).strip() if rest else inline
                if cand and not (cand.startswith("(") and cand.endswith(")")):
                    customer = cand

    items, vat_pct, file_totals = [], None, {}
    for r in range(header_idx + 1, len(rows)):
        row = rows[r]
        texts = [_norm(v) for v in row if isinstance(v, str)]
        joined = " ".join(texts)
        if joined.startswith("tong cong"):
            if "truoc thue" in joined:
                file_totals = {"kg": _num(get(row, "total_kg")), "amount": _num(get(row, "amount"))}
            continue
        raw = " ".join(str(v) for v in row if v is not None)
        m = re.search(r"(\d+(?:[.,]\d+)?)\s*%", raw)
        if re.search(r"\b(thue|vat)\b", joined) and m and _num(get(row, "qty")) is None:
            vat_pct = float(m.group(1).replace(",", "."))
            continue
        name = get(row, "name")
        name = str(name).strip() if name is not None else ""
        qty, per, total = _num(get(row, "qty")), _num(get(row, "kg_per_unit")), _num(get(row, "total_kg"))
        price = _num(get(row, "price"))
        if not name and qty is None and total is None:
            continue  # dòng trống
        if not name:
            warnings.append(f"Dòng {r + 1}: thiếu tên hàng hóa — bỏ qua")
            continue
        line = f"Dòng {r + 1} ({name[:40]})"
        if not qty:
            if per and total:
                qty = round(total / per, 3)
            else:
                warnings.append(f"{line}: thiếu số lượng — bỏ qua")
                continue
        if per:
            kg = qty * per
            if total is not None and abs(total - kg) > 0.5:
                warnings.append(f"{line}: Tổng KL trong file {total:g} kg ≠ SL × KL/1 bộ = {kg:g} kg — lấy theo SL × KL/1 bộ")
        elif total:
            kg, per = total, round(total / qty, 3)
        else:
            warnings.append(f"{line}: thiếu KL/1 bộ và Tổng KL — bỏ qua")
            continue
        if not price:
            warnings.append(f"{line}: chưa có đơn giá")
            price = 0
        amount = round(kg * price)
        file_amount = _num(get(row, "amount"))
        if file_amount is not None and abs(file_amount - amount) > 1:
            warnings.append(f"{line}: Thành tiền trong file {file_amount:,.0f} ≠ Tổng KL × Đơn giá = {amount:,.0f}")
        unit = get(row, "unit")
        note = get(row, "note")
        items.append({"name": name, "unit": str(unit).strip() if unit else "Bộ", "qty": qty, "kgPerUnit": per,
                      "kg": round(kg, 3), "price": price, "amount": amount,
                      "note": str(note).strip() if note is not None else ""})

    if not items:
        raise HTTPException(400, "File không có dòng hàng hóa nào hợp lệ")
    total_kg = round(sum(i["kg"] for i in items), 3)
    value = round(sum(i["kg"] * i["price"] for i in items))
    if vat_pct is None:
        vat_pct = 10
        warnings.append("Không thấy dòng THUẾ VAT trong file — tạm lấy 10%")
    if file_totals.get("amount") is not None and abs(file_totals["amount"] - value) > 1:
        warnings.append(f"Tổng cộng trước thuế trong file {file_totals['amount']:,.0f} ≠ tổng tính lại {value:,.0f}")
    vat = round(value * vat_pct / 100)
    return {"sheet": sheet, "customer": customer, "items": items, "vatPct": vat_pct, "totalKg": total_kg,
            "value": value, "vatAmount": vat, "valueAfterVat": value + vat, "warnings": warnings}


# ---------------------------------------------------------------- xuất ra đúng mẫu
def order_excel(o=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Đặt hàng"
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap = Alignment(vertical="center", wrap_text=True)

    ws.merge_cells("A1:I1")
    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = center
    ws.merge_cells("A2:B2")
    ws.merge_cells("C2:I2")
    ws["A2"] = "CÔNG TY ĐẶT HÀNG:"
    ws["A2"].font = bold
    ws["C2"] = o.customer if o else CUSTOMER_PLACEHOLDER
    ws["C2"].font = bold
    for c, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font, cell.alignment, cell.border = bold, center, box
        cell.fill = PatternFill("solid", fgColor="E7E2D6")

    items = list(o.items) if o else []
    n = max(len(items), 1 if o else 6)
    first, last = 4, 4 + n - 1
    for k in range(n):
        r = first + k
        i = items[k] if k < len(items) else None
        per = (i.kg_per_unit if i.kg_per_unit is not None else (i.kg / i.qty if i.qty else 0)) if i else None
        vals = [k + 1, i.name if i else None, i.unit if i else None, i.qty if i else None, per,
                f"=E{r}*D{r}", i.price if i else None, f"=G{r}*F{r}", (i.note or None) if i else None]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = box
            cell.alignment = wrap if c in (2, 9) else Alignment(horizontal="center" if c in (1, 3) else "right", vertical="center")
            if c in (4, 5, 6, 7, 8):
                cell.number_format = "#,##0.###" if c in (4, 5, 6) else "#,##0"

    vat = o.vat_pct if o and o.vat_pct is not None else 10
    t1, t2, t3 = last + 1, last + 2, last + 3
    for r, label, formula in ((t1, "TỔNG CỘNG TRƯỚC THUẾ", f"=SUM(H{first}:H{last})"),
                              (t2, f"THUẾ VAT {vat:g}%", f"=+H{t1}*{vat:g}%"),
                              (t3, "TỔNG CỘNG SAU THUẾ", f"=+H{t2}+H{t1}")):
        ws.merge_cells(f"A{r}:E{r}")
        ws[f"A{r}"] = label
        ws[f"A{r}"].font, ws[f"A{r}"].alignment = bold, center
        ws[f"H{r}"] = formula
        ws[f"H{r}"].number_format = "#,##0"
        ws[f"H{r}"].font = bold
        for c in range(1, 10):
            ws.cell(row=r, column=c).border = box
    ws[f"F{t1}"] = f"=SUM(F{first}:F{last})"
    ws[f"F{t1}"].number_format = "#,##0.###"
    ws[f"F{t1}"].font = bold

    for c, w in enumerate([6, 46, 8, 11, 10, 14, 12, 16, 20], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A4"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
