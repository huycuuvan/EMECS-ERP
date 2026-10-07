"""Xuất Excel thật (.xlsx, openpyxl) cho các sổ danh sách.

Mỗi sổ: dòng tiêu đề (tên sổ) + dòng thông tin (giờ xuất, người xuất, bộ lọc) + dòng tiêu đề cột in đậm, cố định
khi cuộn, có lọc tự động; định dạng số kg `#,##0`, tiền `#,##0 "₫"`, ngày giờ `dd/mm/yyyy hh:mm` theo giờ Việt Nam;
độ rộng cột tự tính theo nội dung. Bộ lọc nhận giống API danh sách + `q` (tìm chữ) + `ids` (mã đang hiển thị).
"""
from collections import defaultdict
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import services as svc
from .config import PC_FILL_HOURS
from .db import utcnow
from .models import Contract, Lsx, Mismatch, Order, Receipt, Task, User, VLoss, Weighing
from .utils import VN_TZ, fmt_num

FMT = {
    "kg": "#,##0", "num": "#,##0", "money": '#,##0 "₫"', "pct": '0"%"', "delta": "+#,##0;-#,##0;0",
    "dt": "dd/mm/yyyy hh:mm", "date": "dd/mm/yyyy",
}
STATUS_TASK_TYPE = {"di_ma": "Chở hàng đi mạ", "giao_khach": "Lấy mạ → giao khách"}


@dataclass
class Filters:
    contract_id: str | None = None
    lsx_id: str | None = None
    status: str | None = None
    type: str | None = None
    driver: str | None = None
    source: str | None = None
    due: str | None = None
    kind: str | None = None
    date_from: datetime | None = None
    date_to: datetime | None = None
    q: str | None = None
    ids: list[str] = field(default_factory=list)

    def describe(self) -> str:
        parts = []
        if self.contract_id:
            parts.append(f"HĐ {self.contract_id}")
        if self.lsx_id:
            parts.append(f"LSX {self.lsx_id}")
        if self.status:
            parts.append(f"trạng thái: {self.status}")
        if self.type:
            parts.append(f"loại: {STATUS_TASK_TYPE.get(self.type, self.type)}")
        if self.driver:
            parts.append(f"tài xế: {self.driver}")
        if self.source:
            parts.append(f"nguồn: {self.source}")
        if self.due:
            parts.append(f"ngày hoàn thành: {self.due}")
        if self.kind:
            parts.append(f"loại phiếu: {self.kind}")
        if self.date_from or self.date_to:
            f = lambda d: _local(d).strftime("%d/%m/%Y") if d else "…"
            parts.append(f"kỳ {f(self.date_from)} – {f(self.date_to)}")
        if self.q:
            parts.append(f'tìm "{self.q}"')
        if self.ids:
            parts.append(f"{len(self.ids)} dòng đang lọc trên màn hình")
        return "; ".join(parts)


def _local(d: datetime | None) -> datetime | None:
    """UTC → giờ Việt Nam, bỏ tz (Excel không lưu múi giờ)."""
    if d is None:
        return None
    if d.tzinfo is None:
        d = d.replace(tzinfo=timezone.utc)
    return d.astimezone(VN_TZ).replace(tzinfo=None)


Row = tuple[str, list[Any]]
Loader = Callable[[Session, Filters, User], list[Row]]


@dataclass
class Sheet:
    title: str
    page: str  # trang phân quyền (security.can)
    file: str  # tên file (không dấu)
    cols: list[tuple[str, str]]  # (tiêu đề cột, kiểu: text|kg|num|money|pct|delta|dt|date)
    rows: Loader


def _eq(q, col, v):
    return q.where(col == v) if v else q


# ---------------------------------------------------------------- nạp dữ liệu từng sổ
def _orders(db: Session, f: Filters, _u: User) -> list[Row]:
    q = _eq(select(Order), Order.status, f.status).order_by(Order.date.desc())
    out = []
    for o in db.scalars(q):
        items = "; ".join(f"{i.name} ({fmt_num(i.qty)} {i.unit}, {fmt_num(i.kg)} kg, {fmt_num(i.price)} ₫/kg)"
                          for i in o.items)
        out.append((o.id, [o.id, o.code, o.customer, o.date, o.file or "", items, o.total_kg, o.value, o.status,
                           o.contract_id or "", o.note]))
    return out


def _contracts(db: Session, f: Filters, _u: User) -> list[Row]:
    giao = defaultdict(list)
    for t in db.scalars(select(Task).where(Task.type == "giao_khach", Task.kg_delivered.is_not(None))):
        giao[t.contract_id].append(t.kg_delivered)
    q = _eq(select(Contract), Contract.status, f.status).order_by(Contract.sent_to_kt_at.desc())
    out = []
    for c in db.scalars(q):
        adv = svc.advance_info(c)
        delivered = sum(giao[c.id])
        comp = svc.complete_info(c)
        if f.due and comp["state"] != f.due:
            continue
        b = svc.billing(db, c)  # công nợ: kg khách ký nhận × đơn giá, không VAT
        dval, paid = b["deliveredValue"], b["paidTotal"]
        out.append((c.id, [c.id, c.order_id, c.code, c.customer, c.total_qty, c.unit, c.total_kg, c.unit_price,
                           c.value, c.sent_to_kt_at, c.complete_by, comp["label"], c.deliver_by, svc.deliver_info(c, delivered)["label"], c.sign_date or "Chưa nhận về", c.advance_pct,
                           c.advance_required, c.advance_received, adv["label"], delivered, len(giao[c.id]), dval,
                           paid, dval - paid, c.status, c.owner]))
    return out


def _lsx(db: Session, f: Filters, _u: User) -> list[Row]:
    customers = {c.id: c.customer for c in db.scalars(select(Contract))}
    q = _eq(select(Lsx), Lsx.contract_id, f.contract_id)
    if f.status and f.status != "late":
        q = q.where(Lsx.status == f.status)
    now = utcnow()
    out = []
    for x in db.scalars(q.order_by(Lsx.assigned_at.desc())):
        if f.status == "late" and not (x.status not in ("Hoàn thành", "Từ chối") and (x.ext_to or x.deadline) < now):
            continue
        pct = round(x.qty_done / x.qty_plan * 100) if x.qty_plan else 0
        out.append((x.id, [x.id, x.name, x.contract_id, customers.get(x.contract_id, ""), x.assigned_at, x.assigned_by,
                           x.lead_days, x.deadline, x.ext_to or "", x.qty_plan, x.kg_plan, x.qty_done, x.kg_done, pct,
                           x.status, x.accepted_at, x.accepted_by or "", x.reject_reason or ""]))
    return out


def _receipts(db: Session, f: Filters, _u: User) -> list[Row]:
    q = _eq(_eq(select(Receipt), Receipt.contract_id, f.contract_id), Receipt.lsx_id, f.lsx_id)
    return [(r.id, [r.id, r.lsx_id, r.contract_id, r.date, r.qty, r.kg, r.by, r.note])
            for r in db.scalars(q.order_by(Receipt.date.desc()))]


def _weighings(db: Session, f: Filters, _u: User) -> list[Row]:
    q = _eq(select(Weighing), Weighing.contract_id, f.contract_id)
    if f.status and f.status not in ("miss", "overdue"):
        q = q.where(Weighing.status == f.status)
    now = utcnow()
    out = []
    for p in db.scalars(q.order_by(Weighing.date.desc())):
        miss = p.kg_actual is None or not p.photo
        if f.status == "miss" and not miss:
            continue
        if f.status == "overdue" and not (miss and (now - p.date).total_seconds() / 3600 >= PC_FILL_HOURS):
            continue
        delta = p.kg_actual - p.kg_expected if p.kg_actual is not None else None
        out.append((p.id, [p.id, p.date, p.contract_id, p.lsx_id, p.kg_expected,
                           p.kg_actual if p.kg_actual is not None else "Chưa cân", delta,
                           "Đã có ảnh" if p.photo else "Thiếu ảnh", p.signer_boc_xep, p.signer_kho, p.signer_lai_xe,
                           p.by, p.status, p.mismatch_id or ""]))
    return out


def _task_overdue(t: Task, now: datetime) -> bool:
    if t.status in ("Từ chối", "Chờ xác nhận") or not t.fill_deadline:
        return False
    missing = (t.kg_at_galv is None or not t.photo) if t.type == "di_ma" else (t.kg_delivered is None or not t.photo)
    return missing and now > t.fill_deadline


def _tasks(db: Session, f: Filters, user: User) -> list[Row]:
    driver = user.name if user.role_list == ["lx"] else f.driver  # lái xe thuần chỉ thấy thẻ của mình
    q = _eq(_eq(_eq(select(Task), Task.contract_id, f.contract_id), Task.type, f.type), Task.driver, driver)
    if f.status and f.status not in ("overdue", "running"):
        q = q.where(Task.status == f.status)
    now = utcnow()
    out = []
    for t in db.scalars(q.order_by(Task.assigned_at.desc())):
        over = _task_overdue(t, now)
        if f.status == "overdue" and not over:
            continue
        if f.status == "running" and (t.status not in ("Đã nhận", "Đang chạy") or over):
            continue
        out.append((t.id, [t.id, STATUS_TASK_TYPE.get(t.type, t.type), t.driver, t.contract_id, t.ref_id or "",
                           t.assigned_at, t.kg_required, t.kg_at_galv, t.kg_picked, t.kg_delivered, t.departed_at,
                           t.fill_deadline, t.filled_at, "Có" if t.photo else "Chưa có",
                           t.status + (" — quá hạn điền" if over else ""), t.reject_reason or "", t.mismatch_id or "",
                           t.note]))
    return out


def _mismatches(db: Session, f: Filters, _u: User) -> list[Row]:
    q = _eq(_eq(select(Mismatch), Mismatch.contract_id, f.contract_id), Mismatch.status, f.status)
    return [(m.id, [m.id, m.source, m.ref_id, m.contract_id, m.expected, m.actual, m.delta, m.reason, m.reason_note,
                    m.reported_by, m.dept, m.date, m.status, m.signed_by or "", m.signed_at])
            for m in db.scalars(q.order_by(Mismatch.date.desc()))]


def _vloss(db: Session, f: Filters, _u: User) -> list[Row]:
    q = _eq(_eq(_eq(select(VLoss), VLoss.status, f.status), VLoss.source, f.source), VLoss.contract_id, f.contract_id)
    return [(e.id, [e.id, e.date, e.source, e.ref_id, e.contract_id, abs(e.kg), e.approved_by, e.note, e.status,
                    e.resolution or "", e.resolved_at, e.resolved_note or ""])
            for e in db.scalars(q.order_by(VLoss.date.desc()))]


def _movement(db: Session, f: Filters, _u: User) -> list[Row]:
    out = []
    for x in svc.movement_log(db, f.contract_id, f.date_from, f.date_to):
        if f.kind and x["kind"] != f.kind:
            continue
        d = datetime.fromisoformat(x["date"]) if x["date"] else None
        out.append((x["id"], [d, x["kind"], x["id"], x["contractId"], x["kg"] if x["kg"] is not None else "chưa điền",
                              x["desc"], x["who"]]))
    return out


SHEETS: dict[str, Sheet] = {
    "orders": Sheet("Sổ đơn hàng khách", "don-hang", "don-hang", [
        ("Mã đơn", "text"), ("Mã nội bộ", "text"), ("Khách hàng", "text"), ("Ngày chốt", "date"),
        ("File ký chốt", "text"), ("Hàng hóa", "text"), ("Khối lượng (kg)", "kg"), ("Giá trị", "money"),
        ("Trạng thái", "text"), ("Hợp đồng", "text"), ("Ghi chú", "text")], _orders),
    "contracts": Sheet("Sổ hợp đồng & tạm ứng", "hop-dong", "hop-dong", [
        ("Mã HĐ", "text"), ("Mã đơn", "text"), ("Mã nội bộ", "text"), ("Khách hàng", "text"), ("Số lượng", "num"),
        ("Đơn vị", "text"), ("Khối lượng (kg)", "kg"), ("Đơn giá (₫/kg)", "money"), ("Giá trị HĐ", "money"),
        ("Chuyển kế toán", "date"), ("Hạn trả HĐ", "date"), ("Tình trạng hạn trả HĐ", "text"), ("Hạn giao hàng", "date"), ("Tình trạng giao hàng", "text"), ("Ngày ký", "date"),
        ("Tạm ứng %", "pct"), ("Tạm ứng yêu cầu", "money"), ("Tạm ứng đã về", "money"), ("Tình trạng tạm ứng", "text"),
        ("Đã giao (kg)", "kg"), ("Số chuyến đã giao", "num"), ("Giá trị hàng đã giao", "money"),
        ("Tiền về lũy kế", "money"), ("Công nợ", "money"), ("Trạng thái", "text"), ("Phụ trách", "text")], _contracts),
    "lsx": Sheet("Sổ lệnh sản xuất", "lsx", "lenh-san-xuat", [
        ("Mã lệnh", "text"), ("Tên lệnh", "text"), ("Hợp đồng", "text"), ("Khách hàng", "text"),
        ("Phát lệnh lúc", "dt"), ("Phát lệnh bởi", "text"), ("Tiến độ (ngày)", "num"), ("Hạn hoàn thành", "date"),
        ("Gia hạn đến", "date"), ("SL kế hoạch", "num"), ("KL kế hoạch (kg)", "kg"), ("SL đã xong", "num"),
        ("KL đã xong (kg)", "kg"), ("% hoàn thành", "pct"), ("Trạng thái", "text"), ("Nhận lệnh lúc", "dt"),
        ("Người nhận lệnh", "text"), ("Lý do từ chối", "text")], _lsx),
    "receipts": Sheet("Sổ chuẩn bị hàng", "tiep-nhan", "phieu-tiep-nhan", [
        ("Mã phiếu", "text"), ("Lệnh SX", "text"), ("Hợp đồng", "text"), ("Ngày giờ", "dt"), ("SL SP", "num"),
        ("Khối lượng (kg)", "kg"), ("Người lập", "text"), ("Ghi chú", "text")], _receipts),
    "weighings": Sheet("Sổ phiếu cân xuất hàng — ký 3 bên", "phieu-can", "phieu-can", [
        ("Mã phiếu", "text"), ("Ngày", "dt"), ("Hợp đồng", "text"), ("Lệnh SX", "text"),
        ("KL theo lệnh xuất (kg)", "kg"), ("KL cân thực (kg)", "kg"), ("Chênh (kg)", "delta"), ("Ảnh phiếu", "text"),
        ("Bốc xếp", "text"), ("Thủ kho", "text"), ("Lái xe", "text"), ("Người lập", "text"), ("Trạng thái", "text"),
        ("Biên bản sai lệch", "text")], _weighings),
    "tasks": Sheet("Sổ thẻ công việc lái xe", "van-chuyen", "the-lai-xe", [
        ("Mã thẻ", "text"), ("Loại việc", "text"), ("Tài xế", "text"), ("Hợp đồng", "text"), ("Chứng từ gốc", "text"),
        ("Giao việc lúc", "dt"), ("KG yêu cầu", "kg"), ("Số cân mạ (kg)", "kg"), ("KG ký với mạ", "kg"),
        ("KG khách ký", "kg"), ("Xuất phát", "dt"), ("Hạn điền (24h)", "dt"), ("Điền phiếu lúc", "dt"),
        ("Ảnh phiếu", "text"), ("Trạng thái", "text"), ("Lý do từ chối", "text"), ("Biên bản sai lệch", "text"),
        ("Ghi chú", "text")], _tasks),
    "mismatches": Sheet("Sổ biên bản sai lệch", "sai-lech", "sai-lech", [
        ("Mã", "text"), ("Điểm phát sinh", "text"), ("Chứng từ gốc", "text"), ("Hợp đồng", "text"),
        ("Kỳ vọng (kg)", "kg"), ("Thực tế (kg)", "kg"), ("Chênh (kg)", "delta"), ("Lý do", "text"),
        ("Diễn giải", "text"), ("Người báo", "text"), ("Phòng", "text"), ("Ngày", "dt"), ("Trạng thái", "text"),
        ("Người ký", "text"), ("Ký lúc", "dt")], _mismatches),
    "vloss": Sheet("Sổ kho ảo chênh lệch", "kho-ao", "kho-ao", [
        ("Bút toán", "text"), ("Ngày", "dt"), ("Nguồn chênh", "text"), ("Phiếu gốc", "text"), ("Hợp đồng", "text"),
        ("KL rơi rớt (kg)", "kg"), ("Người duyệt", "text"), ("Ghi chú", "text"), ("Trạng thái", "text"),
        ("Hướng xử lý", "text"), ("Xử lý lúc", "dt"), ("Ghi chú xử lý", "text")], _vloss),
    "movement-log": Sheet("Nhật ký chứng từ trong kỳ", "bao-cao", "nhat-ky-chung-tu", [
        ("Ngày", "dt"), ("Loại phiếu", "text"), ("Mã chứng từ", "text"), ("Hợp đồng", "text"),
        ("Khối lượng (kg)", "kg"), ("Diễn giải", "text"), ("Người thực hiện", "text")], _movement),
}


# ---------------------------------------------------------------- dựng workbook
def _display_len(v: Any, kind: str) -> int:
    if v is None or v == "":
        return 0
    if isinstance(v, datetime):
        return 16 if kind == "dt" else 10
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        s = f"{abs(v):,.0f}" + (" ₫" if kind == "money" else "") + ("%" if kind == "pct" else "")
        return len(s) + (1 if v < 0 or kind == "delta" else 0)
    return max(len(line) for line in str(v).split("\n"))


def build_xlsx(sheet: Sheet, rows: list[Row], user: User, filters: Filters) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet.title[:31]
    ncol = len(sheet.cols)
    last = get_column_letter(ncol)
    now = utcnow().astimezone(VN_TZ)

    ws["A1"] = sheet.title.upper()
    ws["A1"].font = Font(bold=True, size=14)
    info = f"Xuất lúc {now:%d/%m/%Y %H:%M} · Người xuất: {user.name} · {len(rows)} dòng"
    desc = filters.describe()
    ws["A2"] = info + (f" · Bộ lọc: {desc}" if desc else "")
    ws["A2"].font = Font(italic=True, color="666666")
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")

    head_row = 3
    thin = Side(style="thin", color="BBBBBB")
    widths = [len(h) for h, _ in sheet.cols]
    for j, (h, _) in enumerate(sheet.cols, start=1):
        c = ws.cell(row=head_row, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EDE7DD")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(bottom=thin)

    for i, (_id, values) in enumerate(rows, start=head_row + 1):
        for j, ((_, kind), v) in enumerate(zip(sheet.cols, values), start=1):
            if isinstance(v, datetime):
                v = _local(v)
            c = ws.cell(row=i, column=j, value=v)
            if kind in FMT and (isinstance(v, (int, float, datetime)) and not isinstance(v, bool)):
                c.number_format = FMT[kind]
            if isinstance(v, str) and "\n" in v:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            widths[j - 1] = max(widths[j - 1], _display_len(v, kind))

    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(w + 2, 8), 60)
    ws.row_dimensions[head_row].height = 30
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.auto_filter.ref = f"A{head_row}:{last}{max(head_row, head_row + len(rows))}"
    ws.sheet_view.zoomScale = 100

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export(db: Session, name: str, user: User, f: Filters) -> tuple[bytes, str]:
    sheet = SHEETS[name]
    rows = sheet.rows(db, f, user)
    if f.ids:
        keep = set(f.ids)
        rows = [r for r in rows if r[0] in keep]
    if f.q:
        s = f.q.strip().lower()
        rows = [r for r in rows if s in " ".join("" if v is None else str(v) for v in r[1]).lower()]
    filename = f"{sheet.file}_{utcnow().astimezone(VN_TZ):%Y-%m-%d}.xlsx"
    return build_xlsx(sheet, rows, user, f), filename
