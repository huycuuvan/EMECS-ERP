"""Sửa chứng từ có lịch sử, giao lại thẻ / phát lại lệnh, xuất Excel."""
import os
import tempfile
from io import BytesIO

# giống test_flows: SQLite tạm, hoặc TEST_DATABASE_URL (DB riêng) — không bao giờ dùng DB dev
os.environ["DATABASE_URL"] = os.environ.get("TEST_DATABASE_URL", f"sqlite:///{tempfile.mkdtemp()}/test.db")
os.environ["PBKDF2_ITER"] = "1000"

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.main import app  # noqa: E402

PW = "steel123"
PHONES = {"ql": "0900000001", "kt": "0900000002", "sx": "0900000003", "kho": "0900000004", "lx1": "0900000005",
          "lx2": "0900000006"}


def login(client, who="ql"):
    r = client.post("/api/auth/login", json={"phone": PHONES[who], "password": PW})
    assert r.status_code == 200, r.text
    client.headers["Authorization"] = f"Bearer {r.json()['token']}"
    return r.json()["user"]


@pytest.fixture()
def c():
    with TestClient(app) as client:
        login(client)
        client.post("/api/admin/reset")
        login(client)
        yield client


def hist(c, t, i):
    r = c.get(f"/api/history/{t}/{i}")
    assert r.status_code == 200, r.text
    return r.json()


# ---------------------------------------------------------------- phiếu cân
def test_weighing_kg_edit_creates_mismatch_and_history(c):
    login(c, "kho")
    r = c.patch("/api/weighings/PC-0201", json={"kgActual": 9950})
    assert r.status_code == 400 and "lý do" in r.json()["detail"]  # sửa số kg bắt buộc có lý do
    r = c.patch("/api/weighings/PC-0201", json={"kgActual": 9950, "reason": "Cân lại — đầu cân lệch",
                                                "mismatchReason": "Sai số thiết bị cân"})
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["status"] == "Lệch — chờ ký" and p["mismatchId"]
    assert p["editOutcome"]["mismatchAction"] == "created" and p["editOutcome"]["mismatchId"] == p["mismatchId"]
    mid = p["mismatchId"]
    m = c.get(f"/api/mismatches/{mid}").json()
    assert m["expected"] == 10000 and m["actual"] == 9950 and m["reason"] == "Sai số thiết bị cân"

    h = hist(c, "pc", "PC-0201")
    kg = next(x for x in h if x["field"] == "kgActual")
    assert kg["oldValue"] == "10000" and kg["newValue"] == "9950"
    assert kg["userName"] == "Ngô Minh Kho" and kg["reason"] == "Cân lại — đầu cân lệch"
    assert {"status", "mismatchId"} <= {x["field"] for x in h}

    # vẫn lệch → cập nhật số trên biên bản đang chờ ký (không lập biên bản thứ 2)
    p = c.patch("/api/weighings/PC-0201", json={"kgActual": 9940, "reason": "Nhập nhầm"}).json()
    assert p["editOutcome"]["mismatchAction"] == "updated" and p["mismatchId"] == mid
    m = c.get(f"/api/mismatches/{mid}").json()
    assert m["actual"] == 9940 and m["delta"] == -60
    assert any(x["field"] == "actual" for x in hist(c, "sl", mid))

    # về trong dung sai → giữ biên bản, ghi chú; phiếu vẫn chờ ký
    p = c.patch("/api/weighings/PC-0201", json={"kgActual": 9990, "reason": "Cân lại lần 3"}).json()
    assert p["editOutcome"]["mismatchAction"] == "noted" and p["status"] == "Lệch — chờ ký"
    m = c.get(f"/api/mismatches/{mid}").json()
    assert m["status"] == "Chờ QL ký" and "dung sai" in m["reasonNote"]

    login(c, "ql")
    c.post(f"/api/mismatches/{mid}/sign")
    assert c.get("/api/weighings/PC-0201").json()["status"] == "Đã cân"


def test_weighing_edit_rules(c):
    login(c, "kho")
    # đổi người ký không cần lý do
    p = c.patch("/api/weighings/PC-0201", json={"signers": {"laiXe": "Phạm Văn Tài"}}).json()
    assert p["signers"]["laiXe"] == "Phạm Văn Tài" and p["editOutcome"]["changed"] == ["signers.laiXe"]
    assert c.patch("/api/weighings/PC-0202", json={"kgActual": 1200, "reason": "x"}).status_code == 400  # chưa cân
    assert c.patch("/api/weighings/PC-0003", json={"kgActual": 9000, "reason": "x"}).status_code == 400  # đã vào kho ảo
    p = c.patch("/api/weighings/PC-0201", json={"kgActual": 10010, "reason": "Cân lại"}).json()  # trong dung sai
    assert p["status"] == "Đã cân" and p["mismatchId"] is None and p["editOutcome"]["mismatchAction"] is None
    nochange = c.patch("/api/weighings/PC-0201", json={"kgActual": 10010}).json()
    assert nochange["editOutcome"]["changed"] == []


def test_edit_permissions(c):
    login(c, "kt")
    assert c.patch("/api/weighings/PC-0201", json={"kgActual": 9000, "reason": "x"}).status_code == 403
    assert c.get("/api/history/pc/PC-0201").status_code == 200  # đọc lịch sử: mọi người đăng nhập
    assert c.get("/api/history/hd/HD-2609-01").status_code == 200
    login(c, "lx1")
    assert c.post("/api/tasks/VC-0007/reassign", json={"driver": "Phạm Văn Tài"}).status_code == 403
    assert c.patch("/api/tasks/VC-0007", json={"note": "x"}).status_code == 403
    assert c.patch("/api/mismatches/SL-0002", json={"reasonNote": "x"}).status_code == 403  # không phải người báo
    login(c, "sx")
    assert c.patch("/api/lsx/LSX-HB04A", json={"name": "x"}).status_code == 403
    assert c.post("/api/lsx/LSX-HB04B/reissue", json={}).status_code == 403
    assert c.patch("/api/receipts/PTN-0201", json={"note": "x"}).status_code == 403  # xưởng chỉ xem tiếp nhận


# ---------------------------------------------------------------- giao lại / phát lại
def test_task_reassign_flow(c):
    login(c, "lx2")
    assert c.post("/api/tasks/VC-0007/reject", json={"reason": "Xe đang bảo dưỡng"}).json()["status"] == "Từ chối"
    login(c, "ql")
    assert c.post("/api/tasks/VC-0007/reassign", json={"driver": "Người lạ"}).status_code == 400
    r = c.post("/api/tasks/VC-0007/reassign", json={"driver": "Phạm Văn Tài", "kgRequired": 7800,
                                                    "note": "Đổi xe khác"})
    assert r.status_code == 200, r.text
    t = r.json()
    assert t["status"] == "Chờ xác nhận" and t["driver"] == "Phạm Văn Tài" and t["rejectReason"] is None
    assert t["kgRequired"] == 7800
    h = hist(c, "vc", "VC-0007")
    drv = next(x for x in h if x["field"] == "driver")
    assert drv["oldValue"] == "Lê Đức Vận" and drv["newValue"] == "Phạm Văn Tài" and drv["reason"] == "Đổi xe khác"
    assert c.post("/api/tasks/VC-0001/reassign", json={"driver": "Lê Đức Vận"}).status_code == 400  # đã hoàn thành
    login(c, "lx1")
    assert c.post("/api/tasks/VC-0007/accept").json()["status"] == "Đã nhận"


def test_task_edit_and_kg_correction(c):
    # sửa thông tin trước xuất phát — đổi tài xế thẻ đã nhận → về chờ xác nhận
    t = c.patch("/api/tasks/VC-0105", json={"driver": "Phạm Văn Tài", "note": "đổi người"}).json()
    assert t["status"] == "Chờ xác nhận" and t["acceptedAt"] is None
    assert c.patch("/api/tasks/VC-0105", json={"kgRequired": 9000}).status_code == 400  # thiếu lý do
    assert c.patch("/api/tasks/VC-0001", json={"driver": "Lê Đức Vận"}).status_code == 400  # đã chạy xong
    # sửa số cân mạ đã điền
    assert c.patch("/api/tasks/VC-0003", json={"kgAtGalv": 9900}).status_code == 400
    assert c.patch("/api/tasks/VC-0003", json={"kgPicked": 9900, "reason": "x"}).status_code == 400  # sai loại thẻ
    t = c.patch("/api/tasks/VC-0003", json={"kgAtGalv": 9900, "reason": "Phiếu mạ ghi nhầm"}).json()
    assert t["editOutcome"]["mismatchAction"] == "created" and t["mismatchId"]
    m = c.get(f"/api/mismatches/{t['mismatchId']}").json()
    assert m["source"] == "Cân tại xưởng mạ" and m["delta"] == -80 and m["reportedBy"] == "Lê Đức Vận"
    # giao khách: kg khách ký sửa về khớp kg ký mạ → biên bản đang chờ ký được giữ + ghi chú
    t = c.patch("/api/tasks/VC-0104", json={"kgDelivered": 9940, "reason": "Khách ký bổ sung"}).json()
    assert t["editOutcome"]["mismatchAction"] == "noted" and t["mismatchId"] == "SL-0003"


def test_lsx_reissue_and_edit(c):
    assert c.post("/api/lsx/LSX-HB04A/reissue", json={}).status_code == 400  # chưa bị từ chối
    r = c.post("/api/lsx/LSX-HB04B/reissue", json={"leadDays": 4, "note": "Đã có thép tấm"})
    assert r.status_code == 200, r.text
    x = r.json()
    assert x["status"] == "Chờ nhận" and x["rejectReason"] is None and x["leadDays"] == 4
    assert any("phát lại" in lg["text"] for lg in x["log"])
    assert {"status", "deadline", "rejectReason"} <= {h["field"] for h in hist(c, "lsx", "LSX-HB04B")}
    login(c, "sx")
    assert c.post("/api/lsx/LSX-HB04B/accept").json()["status"] == "Đang SX"

    login(c, "ql")
    assert c.patch("/api/lsx/LSX-HB04A", json={"kgPlan": 16000}).status_code == 400  # thiếu lý do
    x = c.patch("/api/lsx/LSX-HB04A", json={"kgPlan": 16000, "leadDays": 9, "reason": "Khách bổ sung"}).json()
    assert x["kgPlan"] == 16000 and x["leadDays"] == 9 and "sửa lệnh" in x["log"][-1]["text"]
    assert {"kgPlan", "leadDays", "deadline"} <= {h["field"] for h in hist(c, "lsx", "LSX-HB04A")}


def test_receipt_mismatch_order_contract_history(c):
    login(c, "kho")
    assert c.patch("/api/receipts/PTN-0201", json={"kg": 6100}).status_code == 400
    r = c.patch("/api/receipts/PTN-0201", json={"kg": 6100, "note": "Cân lại", "reason": "Đếm sót"}).json()
    assert r["kg"] == 6100 and set(r["editOutcome"]["changed"]) == {"kg", "note"}
    m = c.patch("/api/mismatches/SL-0002", json={"reason": "Sai số thiết bị cân", "reasonNote": "Kiểm lại"}).json()
    assert m["reason"] == "Sai số thiết bị cân"
    assert c.patch("/api/mismatches/SL-0001", json={"reasonNote": "x"}).status_code == 400  # đã ký
    assert c.patch("/api/mismatches/SL-0002", json={"reason": "Lý do bịa"}).status_code == 400
    assert hist(c, "sl", "SL-0002")[0]["userName"] == "Ngô Minh Kho"

    login(c, "ql")
    c.patch("/api/orders/DH-2609-02", json={"note": "Ghi chú mới"})
    assert any(h["field"] == "note" and h["newValue"] == "Ghi chú mới" for h in hist(c, "dh", "DH-2609-02"))
    c.patch("/api/contracts/HD-2609-02", json={"unitPrice": 30000})
    assert {"unitPrice", "value"} <= {h["field"] for h in hist(c, "hd", "HD-2609-02")}


# ---------------------------------------------------------------- xuất Excel
def _wb(r):
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    return load_workbook(BytesIO(r.content))


def test_export_xlsx(c):
    r = c.get("/api/export/weighings.xlsx")
    wb = _wb(r)
    cd = r.headers["content-disposition"]
    assert "phieu-can_" in cd and ".xlsx" in cd and "filename*=UTF-8''" in cd
    ws = wb.active
    assert ws["A3"].value == "Mã phiếu" and ws["A3"].font.b and ws.freeze_panes == "A4"
    assert "Quản lý A" in ws["A2"].value
    n = len(c.get("/api/weighings").json())
    assert ws.max_row - 3 == n
    assert ws["E4"].number_format == "#,##0"

    ws = _wb(c.get("/api/export/weighings.xlsx", params={"contract_id": "HD-2609-06"})).active
    assert [ws.cell(row=i, column=1).value for i in range(4, ws.max_row + 1)] == ["PC-0202", "PC-0201"]
    ws = _wb(c.get("/api/export/contracts.xlsx", params={"ids": "HD-2609-01,HD-2609-02"})).active
    assert ws.max_row == 5 and '"₫"' in ws["I4"].number_format
    ws = _wb(c.get("/api/export/movement-log.xlsx", params={"kind": "Chuẩn bị hàng"})).active
    assert {ws.cell(row=i, column=2).value for i in range(4, ws.max_row + 1)} == {"Chuẩn bị hàng"}
    for kind in ("orders", "lsx", "receipts", "tasks", "mismatches", "vloss"):
        _wb(c.get(f"/api/export/{kind}.xlsx"))
    assert c.get("/api/export/khong-co.xlsx").status_code == 404

    login(c, "kt")
    assert c.get("/api/export/weighings.xlsx").status_code == 403  # kế toán không xem trạm cân
    _wb(c.get("/api/export/contracts.xlsx"))
    login(c, "lx1")  # lái xe chỉ xuất được thẻ của mình
    ws = _wb(c.get("/api/export/tasks.xlsx")).active
    assert {ws.cell(row=i, column=3).value for i in range(4, ws.max_row + 1)} == {"Phạm Văn Tài"}
