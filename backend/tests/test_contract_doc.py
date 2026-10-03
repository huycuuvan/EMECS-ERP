"""Kế toán soạn thảo hợp đồng theo mẫu: tự điền từ đơn hàng + khách hàng, ô KT nhập, xuất Word, 4 bước trạng thái."""
from io import BytesIO

from docx import Document

from test_master import c, login  # noqa: F401

SELLER = {"name": "CÔNG TY TNHH BÊN BÁN MẪU", "address": "Số 1 Đường A, Hà Nội", "phone": "0123 456 789",
          "banks": ["111 ngân hàng X - CN Y", "222 ngân hàng Z"], "taxCode": "0100000001",
          "representative": "Ông Nguyễn Văn B", "title": "Giám đốc"}


def _contract(c):
    cu = c.post("/api/customers", json={
        "name": "CÔNG TY CP KHÁCH MẪU", "address": "Số 2 Đường B, Nam Định", "phone": "0228.000000",
        "taxCode": "0600000000", "bankAccount": "999", "bankName": "Ngân hàng Q - CN Nam Định",
        "representative": "Bà Trần Thị C", "representativeTitle": "Giám đốc", "segment": "Thân thiết"}).json()
    o = c.post("/api/orders", json={"customer": cu["name"], "customerId": cu["id"], "vatPct": 10, "items": [
        {"name": "Xà đỡ thẳng XĐ35-2L", "qty": 15, "unit": "Bộ", "kgPerUnit": 46.5, "price": 28000},
        {"name": "Xà néo XN35-3L", "qty": 14, "unit": "Bộ", "kgPerUnit": 137.4, "price": 28000}]}).json()
    return o, c.post(f"/api/orders/{o['id']}/send-to-kt", json={"completeBy": "2099-06-30"}).json()


def test_seller_settings_admin_only(c):
    assert c.put("/api/settings/seller", json=SELLER).json()["banks"] == SELLER["banks"]
    login(c, "kt")
    assert c.get("/api/settings/seller").json()["name"] == SELLER["name"]
    assert c.put("/api/settings/seller", json=SELLER).status_code == 403


def test_document_autofill_edit_and_docx(c):
    c.put("/api/settings/seller", json=SELLER)
    o, hd = _contract(c)
    d = c.get(f"/api/contracts/{hd['id']}/document").json()
    assert d["number"] == o["id"]  # số HĐ = mã đơn hàng
    assert d["buyer"]["name"] == "CÔNG TY CP KHÁCH MẪU" and d["buyer"]["banks"] == ["999 tại Ngân hàng Q - CN Nam Định"]
    assert d["buyer"]["representative"] == "Bà Trần Thị C"
    assert [(x["name"], x["unit"], x["qty"], x["kg"]) for x in d["lines"]] == \
        [("Xà đỡ thẳng XĐ35-2L", "Bộ", 15, 697.5), ("Xà néo XN35-3L", "Bộ", 14, 1923.6)]
    assert d["lines"][0]["unitPrice"] == 1_302_000 and d["lines"][0]["amount"] == 19_530_000
    assert d["total"] == o["value"] and d["vatPct"] == 10 and d["grandTotal"] == d["total"] + d["vat"]
    assert d["deliveryTime"] == "trước ngày 30/06/2099."  # từ ngày hoàn thành đơn

    login(c, "kt")
    first = d["lines"][0]["itemId"]
    d = c.put(f"/api/contracts/{hd['id']}/document", json={
        "number": "0742/2026/HĐKT", "date": "2026-10-05", "prices": {str(first): 1_302_040},
        "advances": [100_000_000, 50_000_000], "warranty": "18 tháng", "basis": ["Căn cứ nhu cầu hai bên."],
        "buyer": {**d["buyer"], "banks": ["999 tại Ngân hàng Q - CN Nam Định", "888 tại Ngân hàng R"]}}).json()
    assert d["lines"][0]["amount"] == 15 * 1_302_040 and d["lines"][0]["custom"]
    assert d["dateText"] == "ngày 05 tháng 10 năm 2026"
    assert d["words"] == d["wordsAuto"] and d["advances"][0]["words"] == "Một trăm triệu đồng"
    h = c.get(f"/api/contracts/{hd['id']}").json()["contract"]
    assert h["status"] == "Đã soạn thảo" and h["number"] == "0742/2026/HĐKT"
    assert h["value"] == d["total"] and h["advance"]["required"] == 100_000_000

    r = c.get(f"/api/contracts/{hd['id']}/document.docx")
    assert r.status_code == 200 and "wordprocessingml" in r.headers["content-type"]
    doc = Document(BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    for s in ("Số: 0742/2026/HĐKT", "ngày 05 tháng 10 năm 2026", "CÔNG TY CP KHÁCH MẪU", "Tài khoản: 999 tại",
              "Hoặc 888 tại Ngân hàng R", "Đại diện là: Bà Trần Thị C", SELLER["name"], "Tài khoản số: 222 ngân hàng Z",
              "Xà néo XN35-3L", "1.923,6", "1.302.040", "THUẾ VAT 10%", d["words"], "Tạm ứng lần 2: 50.000.000 VND",
              "18 tháng", "Căn cứ nhu cầu hai bên.", "trước ngày 30/06/2099."):
        assert s in text, s
    assert "{{" not in text and "{%" not in text


def test_contract_four_steps_and_lock(c):
    _, hd = _contract(c)
    cid = hd["id"]
    assert c.post(f"/api/contracts/{cid}/returned").status_code == 400  # chưa soạn thảo
    c.put(f"/api/contracts/{cid}/document", json={})
    assert c.post(f"/api/contracts/{cid}/returned").json()["status"] == "Đã gửi khách hàng"
    notes = c.get("/api/notifications").json()
    assert any(f"HĐ {cid} đã gửi khách hàng" in n["title"] for n in notes)  # báo Quản lý
    assert c.post(f"/api/contracts/{cid}/signed").json()["status"] == "Đã nhận về"
    assert c.put(f"/api/contracts/{cid}/document", json={"warranty": "x"}).status_code == 400  # khóa sau khi nhận về
    assert c.post(f"/api/contracts/{cid}/completed").json()["status"] == "Đã hoàn thành"


def test_payment_needs_manager_approval(c):
    login(c, "kt")
    hd = c.post("/api/contracts/HD-2609-04/payments", json={"amount": 480_000_000, "type": "Tạm ứng 30%"}).json()
    pid = hd["payments"][-1]["id"]
    assert hd["advance"]["received"] == 0 and hd["pendingPayment"] == 480_000_000
    assert c.get("/api/contracts/HD-2609-04").json()["paidTotal"] == 0
    assert c.post(f"/api/payments/{pid}/approve").status_code == 403
    login(c, "ql")
    assert [p["id"] for p in c.get("/api/payments/pending").json()] == [pid]
    assert c.get("/api/dashboard").json()["pendingPayments"][0]["amount"] == 480_000_000
    hd = c.post(f"/api/payments/{pid}/approve").json()
    assert hd["advance"]["received"] == 480_000_000 and hd["payments"][-1]["approvedBy"] == "Quản lý A"
    assert c.post(f"/api/payments/{pid}/approve").status_code == 400
    login(c, "kt")
    pid2 = c.post("/api/contracts/HD-2609-04/payments", json={"amount": 1, "type": "Thanh toán"}).json()["payments"][-1]["id"]
    login(c, "ql")
    hd = c.post(f"/api/payments/{pid2}/reject", json={"reason": "Sai số tiền"}).json()
    assert hd["payments"][-1]["status"] == "Từ chối" and c.get("/api/contracts/HD-2609-04").json()["paidTotal"] == 480_000_000
