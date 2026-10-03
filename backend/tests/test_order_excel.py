"""Đơn hàng theo file Excel đặt hàng của khách (BẢNG XÁC NHẬN GIÁ TRỊ VÀ KHỐI LƯỢNG ĐẶT HÀNG)."""
from io import BytesIO

from openpyxl import Workbook, load_workbook

from test_master import c  # noqa: F401  (fixture: client đăng nhập Quản lý, dữ liệu mẫu)


def _customer_file(company="(TỰ LẤY THÔNG TIN KHÁCH HÀNG)") -> bytes:
    """Dựng file giống mẫu khách: công thức chưa có giá trị tính sẵn (như file sinh bằng phần mềm khác)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Đặt hàng"
    ws["A1"] = "BẢNG XÁC NHẬN GIÁ TRỊ VÀ KHỐI LƯỢNG ĐẶT HÀNG"
    ws["A2"], ws["C2"] = "CÔNG TY ĐẶT HÀNG:", company
    ws.append(["STT", "Tên hàng hóa", "ĐVT", "Số lượng", "KL/1 bộ", "Tổng KL (kg)", "Đơn giá", "Thành tiền", "Ghi chú"])
    rows = [("Tiếp địa lặp lại", 790, 20, 27000, None), ("Xà đỡ dây sau công tơ cột đơn XCT-1", 121, 7, 26500, "mạ kẽm"),
            ("Vật tư phụ", 3, 1, 26500, None)]
    for k, (name, qty, per, price, note) in enumerate(rows, start=1):
        r = 3 + k
        ws.append([k, name, "Bộ", qty, per, f"=E{r}*D{r}", price, f"=G{r}*F{r}", note])
    ws.append(["TỔNG CỘNG TRƯỚC THUẾ", None, None, None, None, "=SUM(F4:F6)", None, "=SUM(H4:H6)"])
    ws.append(["THUẾ VAT 10%", None, None, None, None, None, None, "=+H7*10%"])
    ws.append(["TỔNG CỘNG SAU THUẾ", None, None, None, None, None, None, "=+H8+H7"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(c, data: bytes, name="File đơn hàng.xlsx"):
    return c.post("/api/orders/import-excel", files={"file": (name, data, "application/octet-stream")})


def test_import_customer_excel(c):
    r = _upload(c, _customer_file())
    assert r.status_code == 200, r.text
    d = r.json()
    assert [i["name"] for i in d["items"]] == ["Tiếp địa lặp lại", "Xà đỡ dây sau công tơ cột đơn XCT-1", "Vật tư phụ"]
    first = d["items"][0]
    assert (first["qty"], first["kgPerUnit"], first["kg"], first["price"], first["amount"], first["unit"]) == \
        (790, 20, 15800, 27000, 426_600_000, "Bộ")
    assert d["items"][1]["note"] == "mạ kẽm"
    assert d["totalKg"] == 15800 + 847 + 3
    assert d["value"] == 426_600_000 + 22_445_500 + 79_500
    assert d["vatPct"] == 10 and d["vatAmount"] == round(d["value"] * 0.1)
    assert d["valueAfterVat"] == d["value"] + d["vatAmount"]
    assert d["customer"] == "" and d["customerId"] is None and d["warnings"] == []


def test_import_matches_customer_and_rejects_bad_files(c):
    d = _upload(c, _customer_file("Cty CP Kết cấu thép FECON")).json()
    assert d["customer"] == "Cty CP Kết cấu thép FECON" and d["customerId"]
    assert _upload(c, b"khong phai excel").status_code == 400
    assert _upload(c, _customer_file(), "don.xls").status_code == 400


def test_order_kg_per_unit_vat_and_export(c):
    o = c.post("/api/orders", json={"customer": "Cty Test Excel", "vatPct": 10, "items": [
        {"name": "Tiếp địa lặp lại", "qty": 790, "unit": "Bộ", "kgPerUnit": 20, "price": 27000, "note": "đợt 1"},
        {"name": "Thép tấm", "qty": 2, "unit": "tấm", "kg": 500, "price": 20000}]}).json()
    assert o["totalKg"] == 15800 + 500 and o["value"] == 426_600_000 + 10_000_000
    assert o["items"][0]["kg"] == 15800 and o["items"][0]["note"] == "đợt 1" and o["items"][1]["kgPerUnit"] == 250
    assert o["vatPct"] == 10 and o["valueAfterVat"] == round(o["value"] * 1.1)
    hd = c.post(f"/api/orders/{o['id']}/send-to-kt", json={"completeBy": "2099-01-01"}).json()
    assert hd["vatPct"] == 10

    x = c.get(f"/api/orders/{o['id']}/excel")
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    ws = load_workbook(BytesIO(x.content)).active
    assert ws["A1"].value.startswith("BẢNG XÁC NHẬN") and ws["C2"].value == "Cty Test Excel"
    assert [ws.cell(3, k).value for k in range(1, 10)][:6] == ["STT", "Tên hàng hóa", "ĐVT", "Số lượng", "KL/1 bộ", "Tổng KL (kg)"]
    back = _upload(c, x.content, "Don-hang.xlsx").json()  # xuất → đọc lại ra đúng dòng hàng
    assert [(i["name"], i["qty"], i["kg"], i["price"]) for i in back["items"]] == \
        [("Tiếp địa lặp lại", 790, 15800, 27000), ("Thép tấm", 2, 500, 20000)]
    assert back["vatPct"] == 10 and back["customerId"] == o["customerId"]
    assert c.get("/api/orders/excel-template").status_code == 200
