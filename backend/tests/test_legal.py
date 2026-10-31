"""Bộ hồ sơ gửi khách kèm hợp đồng: biên bản giao nhận + chứng chỉ xuất xưởng theo mẫu Word của khách."""
import io
import zipfile

from docx import Document

from test_master import c, login  # noqa: F401


def _text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [cell.text for cell in row.cells]
    return "\n".join(parts)


def _contract(c):
    c.put("/api/settings/seller", json={**c.get("/api/settings/seller").json(), "name": "CÔNG TY TNHH EMECS VIỆT NAM",
                                       "representative": "Trần Ngọc Anh", "title": "Giám đốc", "representativeHonorific": "Bà",
                                       "workshopRep": "Lê Văn Dũng", "workshopTitle": "QLSX", "workshopHonorific": "Ông"})
    o = c.post("/api/orders", json={"customer": "Cty Hồ Sơ", "items": [
        {"name": "Xà đỡ X-TU", "qty": 6, "unit": "Bộ", "kgPerUnit": 50, "price": 20000},
        {"name": "Tiếp địa RC-2", "qty": 269, "unit": "Bộ", "kgPerUnit": 10, "price": 20000}]}).json()
    return c.post(f"/api/orders/{o['id']}/send-to-kt", json={"completeBy": "2099-01-01"}).json()


def test_legal_forms_fill_from_contract(c):
    hd = _contract(c)
    g = c.get(f"/api/contracts/{hd['id']}/legal").json()
    assert [f["key"] for f in g["forms"] if f["ready"]] == ["bien-ban-giao-nhan", "chung-chi-xuat-xuong"]
    assert g["data"]["giverName"] == "Lê Văn Dũng"  # mặc định bên giao = quản đốc / QLSX trong Thông tin công ty

    bb = _text(c.get(f"/api/contracts/{hd['id']}/legal/bien-ban-giao-nhan.docx").content)
    assert "{{" not in bb and "{%" not in bb
    assert "Cty Hồ Sơ" in bb and "Xà đỡ X-TU" in bb and "269" in bb and "Lê Văn Dũng" in bb
    assert "Đại diện là ông/bà" in bb  # chưa có người nhận → để chấm điền tay

    login(c, "kt")
    r = c.put(f"/api/contracts/{hd['id']}/legal", json={"date": "2026-10-16", "project": "Đường dây 375E10.4",
                                                       "receiverName": "Nguyễn Văn Nam", "receiverTitle": "Chỉ huy trưởng",
                                                       "receiverHonorific": "Ông"})
    assert r.status_code == 200 and r.json()["data"]["project"] == "Đường dây 375E10.4"
    bb = _text(c.get(f"/api/contracts/{hd['id']}/legal/bien-ban-giao-nhan.docx").content)
    assert "Hà Nội, ngày 16 tháng 10 năm 2026" in bb and "Đại diện là ông: Nguyễn Văn Nam" in bb
    cc = _text(c.get(f"/api/contracts/{hd['id']}/legal/chung-chi-xuat-xuong.docx").content)
    assert "{{" not in cc and "Đường dây 375E10.4" in cc and "Bà: Trần Ngọc Anh" in cc and "GIÁM ĐỐC" in cc
    assert "Tiếp địa RC-2" in cc

    z = zipfile.ZipFile(io.BytesIO(c.get(f"/api/contracts/{hd['id']}/legal.zip").content))
    names = z.namelist()
    assert len(names) == 3 and any(n.startswith("Hop-dong_") for n in names)
    assert c.get(f"/api/contracts/{hd['id']}/legal/khong-co.docx").status_code == 404


def test_legal_permissions(c):
    hd = _contract(c)
    login(c, "kho")
    assert c.get(f"/api/contracts/{hd['id']}/legal.zip").status_code == 403
    assert c.put(f"/api/contracts/{hd['id']}/legal", json={"project": "x"}).status_code == 403


def _marked(data: bytes) -> list[str]:
    """Đoạn chữ nền vàng trong file Word."""
    d = Document(io.BytesIO(data))
    paras = list(d.paragraphs) + [p for t in d.tables for row in t.rows for cell in row.cells for p in cell.paragraphs]
    return [r.text for p in paras for r in p.runs if r.font.highlight_color]


def test_fields_marked_yellow(c):
    """Ô điền theo hợp đồng nền vàng, nhãn cố định chữ thường; ?mark=false → bản sạch."""
    hd = _contract(c)
    bb = c.get(f"/api/contracts/{hd['id']}/legal/bien-ban-giao-nhan.docx").content
    hl = _marked(bb)
    assert "Cty Hồ Sơ" in hl and "Xà đỡ X-TU" in hl and "Lê Văn Dũng" in hl
    assert not any("BIÊN BẢN" in t or "Chức vụ" in t for t in hl)  # chữ cố định không bôi
    doc = c.get(f"/api/contracts/{hd['id']}/document.docx").content
    hl = _marked(doc)
    assert "Cty Hồ Sơ" in hl and not any(t.startswith("Địa chỉ") for t in hl)  # "Địa chỉ: " đen, phần sau vàng
    assert _marked(c.get(f"/api/contracts/{hd['id']}/document.docx", params={"mark": False}).content) == []
    assert _marked(c.get(f"/api/contracts/{hd['id']}/legal/chung-chi-xuat-xuong.docx?mark=false").content) == []


def test_company_files_in_pack(c):
    hd = _contract(c)
    r = c.post("/api/settings/legal-files", files={"file": ("Đăng ký kinh doanh.pdf", b"%PDF-1.4 dkkd", "application/pdf")})
    assert r.status_code == 200, r.text
    f = r.json()[0]
    assert f["name"] == "Đăng ký kinh doanh.pdf" and f["url"].startswith("/uploads/") and "sig=" in f["url"]
    c.post("/api/settings/legal-files", files={"file": ("Dang ky kinh doanh.pdf", b"%PDF-1.4 ban 2", "application/pdf")})
    assert c.post("/api/settings/legal-files", files={"file": ("virus.exe", b"MZ", "application/octet-stream")}).status_code == 400

    z = zipfile.ZipFile(io.BytesIO(c.get(f"/api/contracts/{hd['id']}/legal.zip").content))
    names = z.namelist()
    assert "Ho-so-cong-ty/Dang-ky-kinh-doanh.pdf" in names and "Ho-so-cong-ty/Dang-ky-kinh-doanh-2.pdf" in names
    assert z.read("Ho-so-cong-ty/Dang-ky-kinh-doanh.pdf") == b"%PDF-1.4 dkkd"

    login(c, "kho")
    assert c.post("/api/settings/legal-files", files={"file": ("a.pdf", b"x", "application/pdf")}).status_code == 403
    login(c)
    for x in c.get("/api/settings/legal-files").json():
        assert c.delete(f"/api/settings/legal-files/{x['id']}").status_code == 200
    assert c.get("/api/settings/legal-files").json() == []
    assert len(zipfile.ZipFile(io.BytesIO(c.get(f"/api/contracts/{hd['id']}/legal.zip").content)).namelist()) == 3


def _png(w: int = 8, h: int = 8) -> bytes:
    """Ảnh PNG đỏ nhỏ (không cần Pillow)."""
    import struct
    import zlib

    def chunk(t: bytes, d: bytes) -> bytes:
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    raw = b"".join(b"\x00" + b"\xdc\x14\x1e" * w for _ in range(h))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _header_images(data: bytes) -> list[int]:
    """Số ảnh trong header từng section của file Word."""
    d = Document(io.BytesIO(data))
    return [len(s.header._element.xpath(".//pic:pic")) for s in d.sections]


def test_company_logo_in_word_header(c):
    """Anh Thắng 31/10: logo công ty tải 1 lần ở Thông tin công ty → in đầu trang hợp đồng + biên bản + chứng chỉ."""
    hd = _contract(c)
    url = f"/api/contracts/{hd['id']}"
    assert _header_images(c.get(f"{url}/document.docx").content) == [0]  # chưa có logo
    png = io.BytesIO(_png())
    login(c, "kt")
    assert c.post("/api/settings/company-logo", files={"file": ("logo.png", png.getvalue(), "image/png")}).status_code == 403
    login(c, "ql")
    assert c.post("/api/settings/company-logo", files={"file": ("logo.gif", b"GIF89a", "image/gif")}).status_code == 400
    s = c.post("/api/settings/company-logo", files={"file": ("logo.png", png.getvalue(), "image/png")}).json()
    assert s["logo"].startswith("/uploads/logo_")
    assert c.get(s["logo"]).status_code == 200
    assert _header_images(c.get(f"{url}/document.docx").content) == [1]
    for key in ("bien-ban-giao-nhan", "chung-chi-xuat-xuong"):
        assert all(n == 1 for n in _header_images(c.get(f"{url}/legal/{key}.docx").content)), key
    z = zipfile.ZipFile(io.BytesIO(c.get(f"{url}/legal.zip").content))
    for n in z.namelist():
        if n.endswith(".docx"):
            assert all(k == 1 for k in _header_images(z.read(n))), n
    # lưu lại thông tin công ty không làm mất logo; xóa logo → file Word không còn logo
    assert c.put("/api/settings/seller", json={**s, "name": "EMECS"}).json()["logo"]
    assert c.delete("/api/settings/company-logo").json()["logo"] is None
    assert _header_images(c.get(f"{url}/document.docx").content) == [0]
